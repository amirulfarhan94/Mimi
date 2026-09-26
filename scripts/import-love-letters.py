#!/usr/bin/env python3
"""Regenerate js/data/loveLetters.js and js/data/specialLetters.js from the review workbook.

The app never reads the workbook; it only uses the generated JS files, so the letters
stay bundled (and offline) with the PWA. You can also edit those JS files directly.

Usage:  pip install openpyxl
        python3 scripts/import-love-letters.py path/to/Mimi_365_Love_Letters_Final_Review.xlsx
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
MONTHS = {m: i for i, m in enumerate(
    ['january', 'february', 'march', 'april', 'may', 'june', 'july',
     'august', 'september', 'october', 'november', 'december'], start=1)}

# Presentation for each occasion (the message text itself comes from the workbook).
OCCASIONS = {
    'Mimi Birthday': dict(id='mimi-birthday', emoji='🎂', theme='birthday',
                          title='Happy Birthday, Sayang ❤️',
                          cardTitle='A Birthday Letter for You',
                          cardSubtitle='Today is all about you… open it ✨'),
    'Amirul Birthday': dict(id='amirul-birthday', emoji='🎂', theme='birthday',
                            title="Someone's Birthday Today! ❤️",
                            cardTitle="Someone's Birthday Today!",
                            cardSubtitle='There’s a little note about it… 🎈'),
    'Anniversary': dict(id='anniversary', emoji='💍', theme='anniversary',
                        title='Happy Anniversary, Sayang ❤️',
                        cardTitle='An Anniversary Letter',
                        cardSubtitle='Something special is waiting for you… 💕'),
}

js = lambda s: json.dumps(s, ensure_ascii=False)


def main(xlsx):
    wb = openpyxl.load_workbook(xlsx, read_only=True)

    rows = list(wb['365 Love Letters'].iter_rows(min_row=2, values_only=True))
    letters = [(int(r[0]), str(r[1]).strip(), str(r[2]).strip()) for r in rows if r and r[2]]
    assert [d for d, _, _ in letters] == list(range(1, 366)), 'expected days 1..365 in order'

    out = [
        '// Mimi Love Letter — 365 daily letters from Amirul, one per calendar day.',
        '// Day 1 = 1 January … Day 365 = 31 December (see js/loveLetter.js for the date mapping).',
        '// Source: Mimi_365_Love_Letters_Final_Review.xlsx (sheet "365 Love Letters").',
        '// Edit a letter by changing its `text`. Keep exactly 365 entries in day order.',
        '// Regenerate from the workbook with: python3 scripts/import-love-letters.py <file.xlsx>',
        '',
        'export const LOVE_LETTERS = [',
    ]
    out += [f'  {{ day: {d}, category: {js(c)}, text: {js(t)} }},' for d, c, t in letters]
    out += ['];', '']
    (ROOT / 'js/data/loveLetters.js').write_text('\n'.join(out), encoding='utf-8')

    specials = []
    for date, occasion, message in wb['Special Dates'].iter_rows(min_row=2, values_only=True):
        if not occasion:
            continue
        day, month = str(date).split()
        text = str(message).strip().replace('\r\n', '\n')
        # A trailing "— Amirul" line is shown as the letter's signature instead of inside the text.
        m = re.search(r'\n\s*[—-]\s*Amirul\s*$', text)
        signature = 'Amirul' if m else None
        if m:
            text = text[:m.start()].rstrip()
        specials.append(dict(month=MONTHS[month.lower()], day=int(day), occasion=occasion,
                             **OCCASIONS[occasion], messages=[dict(text=text, signature=signature)]))

    out = [
        '// Mimi Love Letter — special dates. These override the normal daily letter.',
        '// Source: Mimi_365_Love_Letters_Final_Review.xlsx (sheet "Special Dates").',
        '//',
        '// Each occasion can hold several messages: add more objects to `messages`.',
        '// The popup shows one per year (rotating), and "Another Note" offers the others first.',
        '// `signature: null` shows no "— Amirul" sign-off.',
        '',
        'export const SPECIAL_DATES = [',
    ]
    for s in specials:
        out += [
            '  {',
            f"    id: {js(s['id'])}, month: {s['month']}, day: {s['day']}, occasion: {js(s['occasion'])},",
            f"    emoji: {js(s['emoji'])}, theme: {js(s['theme'])}, title: {js(s['title'])},",
            f"    cardTitle: {js(s['cardTitle'])},",
            f"    cardSubtitle: {js(s['cardSubtitle'])},",
            '    messages: [',
        ]
        for msg in s['messages']:
            out.append(f"      {{ signature: {js(msg['signature'])}, text: {js(msg['text'])} }},")
        out += ['    ],', '  },']
    out += ['];', '']
    (ROOT / 'js/data/specialLetters.js').write_text('\n'.join(out), encoding='utf-8')

    print(f'Wrote {len(letters)} daily letters and {len(specials)} special dates.')


if __name__ == '__main__':
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
